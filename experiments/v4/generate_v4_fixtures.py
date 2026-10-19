from __future__ import annotations

import hashlib
import json
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
BASE = ROOT / "experiments" / "v4"
FIXTURES = BASE / "fixtures"
MANIFESTS = BASE / "manifests"
ASSETS = ROOT / "corpus"
INSERT_NOTE = Path("<LOCAL_WINDOWS_PATH>")
if str(INSERT_NOTE) not in sys.path:
    sys.path.insert(0, str(INSERT_NOTE))
from insert_note import insert_note  # type: ignore  # noqa: E402

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def set_lang(run, language: str) -> None:
    rpr = run._r.get_or_add_rPr()
    node = rpr.find(qn("w:lang"))
    if node is None:
        node = OxmlElement("w:lang")
        rpr.append(node)
    node.set(qn("w:val"), language)


def setup(title: str, language: str = "en-US") -> Document:
    document = Document()
    document.core_properties.title = title
    document.core_properties.language = language
    document.core_properties.author = "A11yPreserve V4 controlled corpus"
    document.core_properties.subject = "Source-grounded preservation contract"
    return document


def add_title(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="Title")
    run = paragraph.add_run(text)
    set_lang(run, "en-US")


def add_body(document: Document, text: str, language: str = "en-US") -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    set_lang(run, language)


def add_heading(document: Document, text: str, level: int) -> None:
    paragraph = document.add_paragraph(style=f"Heading {level}")
    run = paragraph.add_run(text)
    set_lang(run, "en-US")


def add_image(document: Document, filename: str, alt: str, label: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    inline = paragraph.add_run().add_picture(str(ASSETS / filename), width=Inches(1.5))
    inline._inline.docPr.set("descr", alt)
    inline._inline.docPr.set("title", label)


def add_hyperlink(paragraph, text: str, target: str) -> None:
    relation = paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation)
    run = OxmlElement("w:r")
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_header_row(table, row_index: int) -> None:
    props = table.rows[row_index]._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    props.append(marker)


def add_cell_span(cell, kind: str, value: int) -> None:
    props = cell._tc.get_or_add_tcPr()
    node = OxmlElement(f"w:{kind}")
    node.set(qn("w:val"), str(value))
    props.append(node)


def add_numbering(document: Document, fmt: str) -> int:
    numbering = document.part.numbering_part.element
    abstract_ids = [int(item.get(qn("w:abstractNumId"))) for item in numbering.findall(qn("w:abstractNum"))]
    number_ids = [int(item.get(qn("w:numId"))) for item in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids or [0]) + 1
    number_id = max(number_ids or [0]) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    for level in range(3):
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(level))
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), "•" if fmt == "bullet" else f"%{level + 1}.")
        lvl.extend([start, num_fmt, lvl_text])
        abstract.append(lvl)
    numbering.append(abstract)
    number = OxmlElement("w:num")
    number.set(qn("w:numId"), str(number_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    number.append(abstract_ref)
    numbering.append(number)
    return number_id


def set_numbering(paragraph, number_id: int, level: int) -> None:
    props = paragraph._p.get_or_add_pPr()
    numbering = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), str(number_id))
    numbering.extend([ilvl, num_id])
    props.append(numbering)


def add_math(paragraph, tokens: list[str]) -> None:
    math_para = OxmlElement("m:oMathPara")
    math = OxmlElement("m:oMath")
    for token in tokens:
        run = OxmlElement("m:r")
        text = OxmlElement("m:t")
        text.text = token
        run.append(text)
        math.append(run)
    math_para.append(math)
    paragraph._p.append(math_para)


def save_fixture(variant: str, feature: str, document: Document, contract: dict, rows: list[dict]) -> None:
    path = FIXTURES / f"{variant}.docx"
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    manifest = {
        "variant_id": variant,
        "feature_family": feature,
        "source_format": "DOCX",
        "protocol": "experiments/v4/V4_PROTOCOL.md",
        "role": "development" if variant[-1] in "ABC" else "holdout",
        "fixture_path": f"experiments/v4/fixtures/{path.name}",
        "ground_truth": contract,
        "source_sha256": sha256(path),
    }
    manifest_path = MANIFESTS / f"{variant}.json"
    manifest_path.write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    rows.append({
        "variant_id": variant,
        "feature_family": feature,
        "role": manifest["role"],
        "fixture": manifest["fixture_path"],
        "source_sha256": manifest["source_sha256"],
        "manifest": f"experiments/v4/manifests/{manifest_path.name}",
        "manifest_sha256": sha256(manifest_path),
    })


def build_headings(rows: list[dict]) -> None:
    variants = [
        [("Overview", 1), ("Context", 2)],
        [("Overview", 1), ("Context", 2), ("Details", 3)],
        [("A", 1), ("B", 2), ("C", 2), ("D", 2)],
        [("Overview", 1), ("Context", 2), ("Body", 2), ("Detail", 3)],
        [("Part", 1), ("Section", 2), ("Subsection", 3), ("Subdetail", 4)],
    ]
    for index, levels in enumerate(variants):
        d = setup(f"V4 headings {index + 1}")
        add_title(d, f"V4 headings {index + 1}")
        for text, level in levels:
            add_heading(d, text, level)
        add_body(d, "The heading hierarchy is separated from ordinary body text.")
        save_fixture(f"F01_{chr(65 + index)}", "headings", d, {"headings": [{"text": t, "level": l} for t, l in levels]}, rows)


def build_alt_text(rows: list[dict]) -> None:
    values = [
        ["Short chart."],
        ["A long description of a line chart showing enrollment increasing each academic year from 2021 through 2026."],
        ["Punctuation: review — results (2026)!"],
        ["Descripción de la imagen: café, niño y acción."],
        ["First distinct image.", "Second distinct image with separate values."],
    ]
    for index, alts in enumerate(values):
        d = setup(f"V4 alt text {index + 1}")
        add_title(d, f"V4 alt text {index + 1}")
        add_body(d, "Each image description is an author-supplied source value.")
        for n, alt in enumerate(alts, 1):
            add_image(d, "fixture.png" if n % 2 else "decorative.png", alt, f"Alt image {n}")
        save_fixture(f"F02_{chr(65 + index)}", "alt_text", d, {"figures": [{"alt": value, "order": n} for n, value in enumerate(alts, 1)]}, rows)


def build_lists(rows: list[dict]) -> None:
    patterns = [
        [("unordered", 0, "Alpha"), ("unordered", 0, "Beta")],
        [("ordered", 0, "One"), ("ordered", 0, "Two"), ("ordered", 0, "Three")],
        [("unordered", 0, "Parent"), ("unordered", 1, "Child"), ("unordered", 1, "Sibling")],
        [("ordered", 0, "Chapter"), ("unordered", 1, "Topic"), ("ordered", 1, "Subtopic"), ("ordered", 0, "Next chapter")],
        [("ordered", 0, "Restart one"), ("ordered", 0, "Restart two"), ("unordered", 1, "Nested note")],
    ]
    for index, items in enumerate(patterns):
        d = setup(f"V4 lists {index + 1}")
        add_title(d, f"V4 lists {index + 1}")
        ordered = add_numbering(d, "decimal")
        unordered = add_numbering(d, "bullet")
        for text, level, value in items:
            p = d.add_paragraph(value)
            set_numbering(p, ordered if text == "ordered" else unordered, level)
        save_fixture(f"F03_{chr(65 + index)}", "lists", d, {"lists": [{"type": t, "depth": l, "text": v} for t, l, v in items]}, rows)


def build_basic_tables(rows: list[dict]) -> None:
    specs = [
        [["Name", "Value"], ["A", "1"]],
        [["Year", "North", "South"], ["2025", "10", "20"], ["2026", "11", "21"]],
        [["Measure", "Result"], ["Coverage", "82%"], ["Review", "Pass"], ["Notes", "Three rows"]],
        [["Region", "Q1", "Q2", "Q3"], ["North", "1", "2", "3"], ["South", "4", "5", "6"]],
        [["Category", "Count"], ["A", "3"], ["B", "4"]],
    ]
    for index, values in enumerate(specs):
        d = setup(f"V4 basic table {index + 1}")
        add_title(d, f"V4 basic table {index + 1}")
        table = d.add_table(rows=len(values), cols=len(values[0]))
        table.style = "Table Grid"
        for row_index, row in enumerate(values):
            for col_index, value in enumerate(row):
                table.cell(row_index, col_index).text = value
        add_header_row(table, 0)
        save_fixture(f"F04_{chr(65 + index)}", "basic_tables", d, {"dimensions": {"rows": len(values), "columns": len(values[0])}, "header_rows": [0], "cells": values}, rows)


def build_document_language(rows: list[dict]) -> None:
    locales = ["en-US", "en-GB", "es-ES", "fr-FR", "de-DE"]
    for index, locale in enumerate(locales):
        d = setup(f"V4 language {locale}", locale)
        add_title(d, f"V4 document language {locale}")
        add_body(d, f"This document declares the default locale {locale}.", locale)
        save_fixture(f"F05_{chr(65 + index)}", "document_language", d, {"document_language": locale}, rows)


def build_inline_language(rows: list[dict]) -> None:
    specs = [
        [("bonjour", "fr-FR")],
        [("hola mundo", "es-ES")],
        [("Guten Tag, wie geht es Ihnen?", "de-DE")],
        [("bonjour", "fr-FR"), ("hola", "es-ES"), ("ciao", "it-IT")],
        [("English", "en-GB"), ("日本語", "ja-JP")],
    ]
    for index, spans in enumerate(specs):
        d = setup(f"V4 inline language {index + 1}")
        add_title(d, f"V4 inline language {index + 1}")
        p = d.add_paragraph()
        p.add_run("Default text ")
        for text, locale in spans:
            run = p.add_run(text + " ")
            set_lang(run, locale)
        save_fixture(f"F06_{chr(65 + index)}", "inline_language", d, {"document_language": "en-US", "spans": [{"text": t, "language": l} for t, l in spans if l != "en-US"]}, rows)


def build_decorative(rows: list[dict]) -> None:
    specs = [
        [("", True)],
        [("Meaningful chart.", False)],
        [("", True), ("Meaningful chart.", False)],
        [("Meaningful first.", False), ("", True), ("Meaningful last.", False)],
        [("", True), ("", True), ("Decorative texture.", False)],
    ]
    for index, figures in enumerate(specs):
        d = setup(f"V4 decorative state {index + 1}")
        add_title(d, f"V4 decorative state {index + 1}")
        for n, (alt, decorative) in enumerate(figures, 1):
            add_image(d, "decorative.png" if decorative else "fixture.png", alt, f"Figure {n}")
        save_fixture(f"F07_{chr(65 + index)}", "decorative_images", d, {"figures": [{"alt": a, "decorative": dec, "order": n} for n, (a, dec) in enumerate(figures, 1)]}, rows)


def build_links(rows: list[dict]) -> None:
    specs = [
        [("Read the guidance", "https://example.org/guidance")],
        [("https://example.org/raw", "https://example.org/raw")],
        [("Guidance", "https://example.org/guidance"), ("Reference", "https://example.org/reference")],
        [("Same label", "https://example.org/one"), ("Same label", "https://example.org/two")],
        [("Internal section", "#section-2")],
    ]
    for index, links in enumerate(specs):
        d = setup(f"V4 hyperlinks {index + 1}")
        add_title(d, f"V4 hyperlinks {index + 1}")
        p = d.add_paragraph()
        for n, (label, target) in enumerate(links):
            if n:
                p.add_run(" and ")
            add_hyperlink(p, label, target)
        if any(target.startswith("#") for _, target in links):
            add_heading(d, "Section 2", 2)
        save_fixture(f"F08_{chr(65 + index)}", "hyperlinks", d, {"hyperlinks": [{"text": t, "target": u, "order": n} for n, (t, u) in enumerate(links, 1)]}, rows)


def build_metadata(rows: list[dict]) -> None:
    specs = [
        ("Accessible policy guide", "Accessible policy guide"),
        ("Résumé — été 2026", "Résumé — été 2026"),
        ("Title: Findings (draft)", "Title: Findings (draft)"),
        ("A long document title for preservation review", "A long document title for preservation review"),
        ("Visible report title", "Filename-distinct metadata title"),
    ]
    for index, (visible, core) in enumerate(specs):
        d = setup(core)
        d.core_properties.title = core
        d.core_properties.subject = f"Subject {index + 1}: source metadata"
        d.core_properties.keywords = "a11y, preservation, V4"
        add_title(d, visible)
        add_body(d, "Visible title and core metadata are recorded as separate contract values.")
        save_fixture(f"F09_{chr(65 + index)}", "document_metadata", d, {"visible_title": visible, "core_title": core, "subject": d.core_properties.subject, "keywords": d.core_properties.keywords}, rows)


def build_complex_tables(rows: list[dict]) -> None:
    specs = [
        {"span": "column", "rows": 3},
        {"span": "row", "rows": 4},
        {"span": "multi-level", "rows": 4},
        {"span": "groups", "rows": 5},
        {"span": "mixed", "rows": 5},
    ]
    for index, spec in enumerate(specs):
        d = setup(f"V4 complex table {index + 1}")
        add_title(d, f"V4 complex table {index + 1}")
        table = d.add_table(rows=spec["rows"], cols=4)
        table.style = "Table Grid"
        for r in range(spec["rows"]):
            for c in range(4):
                table.cell(r, c).text = f"{spec['span']} {r + 1}.{c + 1}"
        if index in {0, 2, 3, 4}:
            table.cell(0, 0).merge(table.cell(0, 1)).text = f"{spec['span']} group A"
        if index in {1, 4}:
            table.cell(0, 2).merge(table.cell(1, 2)).text = f"{spec['span']} row span"
        add_header_row(table, 0)
        if index >= 2:
            add_header_row(table, 1)
        save_fixture(f"F10_{chr(65 + index)}", "complex_tables", d, {"span_type": spec["span"], "rows": spec["rows"], "columns": 4, "header_rows": [0, 1] if index >= 2 else [0], "span_signature": {"gridSpan": "2" if index in {0, 2, 3, 4} else None, "vMerge": "restart" if index == 4 else ("restart" if index == 1 else None)}}, rows)


def build_footnotes(rows: list[dict]) -> None:
    specs = [
        ["One note body."],
        ["First note.", "Second note."],
        ["Repeated reference note."],
        ["A longer footnote body containing enough text to test relocation and extraction."],
        ["Note in section one.", "Note in section two.", "Third note."]
    ]
    for index, notes in enumerate(specs):
        d = setup(f"V4 footnotes {index + 1}")
        add_title(d, f"V4 footnotes {index + 1}")
        add_body(d, "The source contains a footnote reference for each declared note.")
        temp1 = BASE / f"_F11_{index}_1.docx"
        temp2 = BASE / f"_F11_{index}_2.docx"
        out = FIXTURES / f"F11_{chr(65 + index)}.docx"
        d.save(temp1)
        current = temp1
        for note_index, body_text in enumerate(notes, 1):
            target = temp2 if note_index < len(notes) else out
            insert_note(str(current), str(target), "footnote", f"[[N{note_index}]]", body_text)
            if current != temp1:
                current.unlink(missing_ok=True)
            current = target
            if target == temp2:
                temp2 = BASE / f"_F11_{index}_3.docx"
        temp1.unlink(missing_ok=True)
        for temp in BASE.glob(f"_F11_{index}_*.docx"):
            if temp != out:
                temp.unlink(missing_ok=True)
        manifest = {"notes": [{"id": n, "text": body} for n, body in enumerate(notes, 1)]}
        path = out
        save_manifest = {"variant_id": f"F11_{chr(65 + index)}", "feature_family": "footnotes", "source_format": "DOCX", "protocol": "experiments/v4/V4_PROTOCOL.md", "role": "development" if index < 3 else "holdout", "fixture_path": f"experiments/v4/fixtures/{path.name}", "ground_truth": manifest, "source_sha256": sha256(path)}
        mp = MANIFESTS / f"F11_{chr(65 + index)}.json"
        mp.write_text(json.dumps(save_manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
        rows.append({"variant_id": save_manifest["variant_id"], "feature_family": "footnotes", "role": save_manifest["role"], "fixture": save_manifest["fixture_path"], "source_sha256": save_manifest["source_sha256"], "manifest": f"experiments/v4/manifests/{mp.name}", "manifest_sha256": sha256(mp)})


def build_equations(rows: list[dict]) -> None:
    specs = [["x + y = z"], ["a/b"], ["x^2 + y^2"], ["sqrt(x) = y"], ["E = mc^2", "F = ma"]]
    token_specs = [[list("x+y=z")], [["a", "/", "b"]], [["x", "2", "+", "y", "2"]], [["√", "x", "=", "y"]], [["E", "=", "m", "c", "2"], ["F", "=", "m", "a"]]]
    for index, expressions in enumerate(specs):
        d = setup(f"V4 equations {index + 1}")
        add_title(d, f"V4 equations {index + 1}")
        equations = token_specs[index]
        for tokens in equations:
            p = d.add_paragraph()
            add_math(p, tokens)
        save_fixture(f"F12_{chr(65 + index)}", "equations", d, {"equations": [{"expression": expr, "tokens": tokens} for expr, tokens in zip(expressions, equations)]}, rows)


def build_captions(rows: list[dict]) -> None:
    specs = [
        ["Figure 1. Workflow."],
        ["Figure 1. First.", "Figure 2. Second."],
        ["Figure A. Caption before.", "Figure B. Caption after."],
        ["Figure 1. Unicode café — results."],
        ["Figure 1. One.", "Figure 2. Two.", "Figure 3. Three."]
    ]
    for index, captions in enumerate(specs):
        d = setup(f"V4 captions {index + 1}")
        add_title(d, f"V4 captions {index + 1}")
        for n, caption in enumerate(captions, 1):
            if index == 2 and n == 1:
                p = d.add_paragraph(style="Caption")
                p.add_run(caption)
                add_image(d, "fixture.png", f"Description for figure {n}.", f"Figure {n}")
            else:
                add_image(d, "fixture.png", f"Description for figure {n}.", f"Figure {n}")
                p = d.add_paragraph(style="Caption")
                p.add_run(caption)
        save_fixture(f"F13_{chr(65 + index)}", "captions", d, {"pairs": [{"figure_order": n, "caption": caption} for n, caption in enumerate(captions, 1)]}, rows)


def build() -> list[dict]:
    FIXTURES.mkdir(parents=True, exist_ok=True)
    MANIFESTS.mkdir(parents=True, exist_ok=True)
    rows: list[dict] = []
    build_headings(rows)
    build_alt_text(rows)
    build_lists(rows)
    build_basic_tables(rows)
    build_document_language(rows)
    build_inline_language(rows)
    build_decorative(rows)
    build_links(rows)
    build_metadata(rows)
    build_complex_tables(rows)
    build_footnotes(rows)
    build_equations(rows)
    build_captions(rows)
    rows.sort(key=lambda row: row["variant_id"])
    (BASE / "source_manifest.json").write_text(json.dumps({"protocol": "experiments/v4/V4_PROTOCOL.md", "fixture_count": len(rows), "feature_family_count": 13, "variants_per_feature": 5, "records": rows}, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    return rows


if __name__ == "__main__":
    result = build()
    print(json.dumps({"created": len(result), "features": sorted({row["feature_family"] for row in result}), "holdout": sum(row["role"] == "holdout" for row in result)}, indent=2))
