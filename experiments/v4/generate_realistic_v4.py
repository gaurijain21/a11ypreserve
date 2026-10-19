from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
BASE = ROOT / "experiments" / "v4" / "realistic"
sys.path.insert(0, str(ROOT / "experiments" / "v4"))
from generate_v4_fixtures import add_body, add_header_row, add_heading, add_hyperlink, add_image, add_math, add_numbering, add_title, set_lang, set_numbering, sha256, setup  # type: ignore  # noqa: E402
from v4_source_oracle import parse  # type: ignore  # noqa: E402


def build_one(index: int) -> tuple[Path, list[dict]]:
    themes = ["community accessibility guide", "course syllabus", "policy memo", "research handout", "public meeting report", "instructional guide", "service overview", "evaluation brief", "program brochure", "small annual report"]
    theme = themes[index - 1]
    d = setup(f"A11yPreserve realistic document {index}: {theme}")
    d.core_properties.subject = f"Integrated realistic document {index}"
    d.core_properties.keywords = "accessibility, preservation, integrated"
    add_title(d, f"{theme.title()} — 2026")
    add_heading(d, "Purpose and audience", 1)
    add_body(d, f"This multi-page {theme} explains a practical process for readers, authors, and reviewers. It combines several accessibility semantics in a document someone might author rather than concatenating atomic fixtures.")
    p = d.add_paragraph()
    p.add_run("The guide includes a short ")
    run = p.add_run("bonjour")
    set_lang(run, "fr-FR")
    p.add_run(" language example and a descriptive external resource.")
    add_hyperlink(p, "resource page", f"https://example.org/realistic/{index}")
    ordered = add_numbering(d, "decimal")
    unordered = add_numbering(d, "bullet")
    add_heading(d, "Recommended actions", 2)
    for n, value in enumerate(["Review the source contract", "Export through the documented workflow", "Inspect destination evidence"], 1):
        paragraph = d.add_paragraph(value)
        set_numbering(paragraph, ordered if n == 1 else unordered, 0)
    add_heading(d, "Evidence summary", 1)
    table = d.add_table(rows=4, cols=3)
    table.style = "Table Grid"
    values = [["Property", "Status", "Owner"], ["Headings", "Declared", "Author"], ["Links", "Descriptive", "Editor"], ["Tables", "Header row", "Reviewer"]]
    for r, row in enumerate(values):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    add_header_row(table, 0)
    add_image(d, "fixture.png", f"Illustration for the {theme}.", f"{theme} illustration")
    caption = d.add_paragraph(style="Caption")
    caption.add_run(f"Figure {index}. {theme.title()} overview.")
    add_body(d, "A source-grounded evaluation preserves the declared property values as contracts and does not assign a global accessibility score.")
    d.add_page_break()
    add_heading(d, "Implementation notes", 1)
    add_body(d, "The second page records context, exceptions, and a short narrative so that the document exercises coexistence across paragraphs, lists, links, images, metadata, language, tables, and captions.")
    add_heading(d, "Exception handling", 2)
    add_body(d, "If a destination representation is ambiguous, the procedure abstains and records the reason. A parser failure is not treated as proof that source semantics were absent.")
    if index % 2 == 0:
        paragraph = d.add_paragraph("This statement carries a note marker [[N1]].")
        paragraph.add_run()
    if index % 3 == 0:
        p = d.add_paragraph()
        add_math(p, ["E", "=", "m", "c", "2"])
    path = BASE / f"R{index:02d}.docx"
    d.save(path)
    if index % 2 == 0:
        helper = Path("<LOCAL_WINDOWS_PATH>")
        sys.path.insert(0, str(helper))
        from insert_note import insert_note  # type: ignore
        temp = BASE / f"R{index:02d}_temp.docx"
        insert_note(str(path), str(temp), "footnote", "[[N1]]", f"Supporting note for integrated document {index}.")
        path.unlink()
        temp.rename(path)
    contracts = [
        {"property": "headings", "expected_count": 5, "source_evidence": "word/document.xml paragraph styles"},
        {"property": "inline_language", "expected_language": "fr-FR", "source_evidence": "word/document.xml w:lang"},
        {"property": "hyperlinks", "expected_target": f"https://example.org/realistic/{index}", "source_evidence": "document relationship"},
        {"property": "basic_tables", "header_rows": [0], "source_evidence": "w:tblHeader"},
        {"property": "alternative_text", "expected_alt": f"Illustration for the {theme}.", "source_evidence": "drawing docPr descr"},
        {"property": "captions", "expected_text": f"Figure {index}. {theme.title()} overview.", "source_evidence": "Caption paragraph"},
        {"property": "document_title", "expected_title": f"A11yPreserve realistic document {index}: {theme}", "source_evidence": "core properties"},
    ]
    if index % 2 == 0:
        contracts.append({"property": "footnotes", "expected_count": 1, "source_evidence": "word/footnotes.xml"})
    if index % 3 == 0:
        contracts.append({"property": "equations", "expected_count": 1, "source_evidence": "OMML"})
    return path, contracts


def main() -> None:
    BASE.mkdir(parents=True, exist_ok=True)
    records = []
    for index in range(1, 11):
        path, contracts = build_one(index)
        raw = parse(path)
        records.append({"document_id": path.stem, "path": str(path.relative_to(ROOT)).replace("\\", "/"), "source_sha256": sha256(path), "pages_target": 2, "provenance": "project-created; no external source or redistribution rights asserted", "contracts": contracts, "raw_source_observation": {"headings": len(raw["headings"]), "figures": len(raw["figures"]), "lists": len(raw["lists"]), "tables": len(raw["tables"]), "hyperlinks": len(raw["hyperlinks"]), "footnotes": len(raw["footnotes"]), "equations": len(raw["equations"]), "captions": len(raw["captions"])}})
    report = {"corpus": "V4 realistic/integrated evaluation", "document_count": len(records), "external_documents": 0, "project_created_documents": len(records), "records": records, "scoring": "per source property contract; no global accessibility score; not merged with atomic denominator"}
    (BASE / "REALISTIC_CORPUS_MANIFEST.json").write_text(json.dumps(report, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    (BASE / "REALISTIC_CONTRACTS.json").write_text(json.dumps({"documents": [{"document_id": r["document_id"], "contracts": r["contracts"]} for r in records]}, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps({"documents": len(records), "property_contracts": sum(len(r["contracts"]) for r in records), "external": 0}, indent=2))


if __name__ == "__main__":
    main()
