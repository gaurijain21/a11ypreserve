from __future__ import annotations

import hashlib
import json
import os
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / "corpus" / "v3_variant_set_b"
FIXTURES = BASE / "fixtures"
MANIFESTS = BASE / "manifests"
ASSETS = ROOT / "corpus"
HELPER = os.environ.get("A11YPRESERVE_DOCUMENT_SKILLS")
if HELPER and HELPER not in sys.path:
    sys.path.insert(0, HELPER)
from insert_note import insert_note  # noqa: E402

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"

def q(tag: str) -> str:
    return qn(tag)

def lang(run, value: str) -> None:
    rpr = run._r.get_or_add_rPr()
    node = rpr.find(q("w:lang"))
    if node is None:
        node = OxmlElement("w:lang")
        rpr.append(node)
    node.set(q("w:val"), value)

def setup(title: str, language: str = "en-US") -> Document:
    d = Document()
    d.core_properties.title = title
    d.core_properties.language = language
    d.core_properties.author = "A11yPreserve variant set"
    styles = d.styles.element
    defaults = styles.find(q("w:docDefaults"))
    if defaults is None:
        defaults = OxmlElement("w:docDefaults")
        styles.insert(0, defaults)
    rpr_default = defaults.find(q("w:rPrDefault")) or OxmlElement("w:rPrDefault")
    if rpr_default.getparent() is None:
        defaults.append(rpr_default)
    rpr = rpr_default.find(q("w:rPr")) or OxmlElement("w:rPr")
    if rpr.getparent() is None:
        rpr_default.append(rpr)
    l = rpr.find(q("w:lang")) or OxmlElement("w:lang")
    if l.getparent() is None:
        rpr.append(l)
    l.set(q("w:val"), language)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4", "Caption"):
        d.styles[name].font.name = "Aptos"
        d.styles[name].font.size = Pt(11)
        d.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    return d

def title(d: Document, text: str) -> None:
    p = d.add_paragraph(style="Title")
    r = p.add_run(text); lang(r, "en-US")

def body(d: Document, text: str, language: str = "en-US") -> None:
    p = d.add_paragraph(); r = p.add_run(text); lang(r, language)

def heading(d: Document, text: str, level: int) -> None:
    p = d.add_paragraph(style=f"Heading {level}"); r = p.add_run(text); lang(r, "en-US")

def alt_image(d: Document, filename: str, alt: str, label: str) -> None:
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run().add_picture(str(ASSETS / filename), width=Inches(1.6))
    run._inline.docPr.set("descr", alt); run._inline.docPr.set("title", label)

def hyperlink(p, text: str, target: str) -> None:
    rid = p.part.relate_to(target, RT.HYPERLINK, is_external=True)
    node = OxmlElement("w:hyperlink"); node.set(q("r:id"), rid)
    run = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = text; run.append(t); node.append(run); p._p.append(node)

def header_row(table, row: int) -> None:
    trpr = table.rows[row]._tr.get_or_add_trPr(); n = OxmlElement("w:tblHeader"); n.set(q("w:val"), "true"); trpr.append(n)

def math_expr(p, tokens: list[str]) -> None:
    para = OxmlElement("m:oMathPara"); math = OxmlElement("m:oMath")
    for token in tokens:
        r = OxmlElement("m:r"); t = OxmlElement("m:t"); t.text = token; r.append(t); math.append(r)
    para.append(math); p._p.append(para)

def create_numbering(d: Document, fmt: str) -> int:
    numbering = d.part.numbering_part.element
    aids = [int(x.get(q("w:abstractNumId"))) for x in numbering.findall(q("w:abstractNum"))]
    nids = [int(x.get(q("w:numId"))) for x in numbering.findall(q("w:num"))]
    aid, nid = max(aids or [0]) + 1, max(nids or [0]) + 1
    abstract = OxmlElement("w:abstractNum"); abstract.set(q("w:abstractNumId"), str(aid))
    for level in range(3):
        lvl = OxmlElement("w:lvl"); lvl.set(q("w:ilvl"), str(level))
        start = OxmlElement("w:start"); start.set(q("w:val"), "1"); lvl.append(start)
        nf = OxmlElement("w:numFmt"); nf.set(q("w:val"), fmt); lvl.append(nf)
        lt = OxmlElement("w:lvlText"); lt.set(q("w:val"), "•" if fmt == "bullet" else f"%{level + 1}."); lvl.append(lt)
        abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num"); num.set(q("w:numId"), str(nid)); ref = OxmlElement("w:abstractNumId"); ref.set(q("w:val"), str(aid)); num.append(ref); numbering.append(num)
    return nid

def set_numbering(p, nid: int, level: int) -> None:
    ppr = p._p.get_or_add_pPr(); np = OxmlElement("w:numPr"); il = OxmlElement("w:ilvl"); il.set(q("w:val"), str(level)); ni = OxmlElement("w:numId"); ni.set(q("w:val"), str(nid)); np.extend([il, ni]); ppr.append(np)

def build() -> list[dict]:
    rows = []
    def save(variant: str, feature: str, d: Document, contract: dict):
        path = FIXTURES / f"{variant}.docx"; d.save(path)
        manifest = {"variant_id": variant, "feature_family": feature, "source_format": "DOCX", "contract_version": "v3-variant-set-b", "fixture_path": f"corpus/v3_variant_set_b/fixtures/{path.name}", "ground_truth": contract}
        mp = MANIFESTS / f"{variant}.json"; mp.write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
        rows.append({"variant_id": variant, "feature_family": feature, "fixture": f"corpus/v3_variant_set_b/fixtures/{path.name}", "source_sha256": sha(path), "manifest": f"corpus/v3_variant_set_b/manifests/{mp.name}", "manifest_sha256": sha(mp)})

    d = setup("Variant B01 headings"); title(d, "Variant B01 headings"); heading(d, "Overview", 1); heading(d, "Context", 2); heading(d, "Detail", 3); heading(d, "Subdetail", 4); body(d, "The title is separate from the four-level hierarchy."); save("B01", "headings", d, {"headings": [{"text":"Overview","level":1},{"text":"Context","level":2},{"text":"Detail","level":3},{"text":"Subdetail","level":4}]})
    d = setup("Variant B02 alternative text"); title(d, "Variant B02 alternative text"); body(d, "Two informative images use distinct author-supplied descriptions."); alt_image(d, "fixture.png", "Workflow: source -> pipeline -> PDF.", "Workflow image"); alt_image(d, "decorative.png", "Second figure, 2026 results.", "Results image"); save("B02", "alt_text", d, {"figures":[{"alt":"Workflow: source -> pipeline -> PDF.","order":1},{"alt":"Second figure, 2026 results.","order":2}]})
    d = setup("Variant B03 lists"); title(d, "Variant B03 lists"); number_id=create_numbering(d,"decimal"); bullet_id=create_numbering(d,"bullet")
    for text, level, ordered in [("First",0,True),("Nested alpha",1,False),("Nested beta",1,False),("Second",0,True),("Nested gamma",1,False)]:
        p=d.add_paragraph(); p.add_run(text); set_numbering(p, number_id if ordered else bullet_id, level)
    save("B03", "lists", d, {"lists":[{"type":"ordered","depth":0,"text":"First"},{"type":"unordered","depth":1,"text":"Nested alpha"},{"type":"unordered","depth":1,"text":"Nested beta"},{"type":"ordered","depth":0,"text":"Second"},{"type":"unordered","depth":1,"text":"Nested gamma"}]})
    d = setup("Variant B04 table"); title(d, "Variant B04 table"); t=d.add_table(rows=4, cols=4); t.style="Table Grid"; vals=[["Quarter","Region","Cases handled","Completion rate"],["Q1","North","120","82%"],["Q2","South","240","87%"],["Q3","West","360","91%"]]
    for i,row in enumerate(vals):
        for j,v in enumerate(row): t.cell(i,j).text=v
    header_row(t,0); save("B04", "table", d, {"dimensions":{"rows":4,"columns":4},"header_rows":[0],"cells":vals})
    d = setup("Variante B05 langue", "fr-FR"); title(d, "Variante B05 langue"); body(d, "Cette note décrit la langue par défaut du document.", "fr-FR"); save("B05", "document_language", d, {"document_language":"fr-FR"})
    d = setup("Variant B06 inline language"); title(d, "Variant B06 inline language"); p=d.add_paragraph(); r=p.add_run("English "); lang(r,"en-US"); r=p.add_run("bonjour"); lang(r,"fr-FR"); r=p.add_run(" hola"); lang(r,"es-ES"); save("B06", "inline_language", d, {"document_language":"en-US","runs":[{"text":"bonjour","language":"fr-FR"},{"text":" hola","language":"es-ES"}]})
    d = setup("Variant B07 decorative state"); title(d, "Variant B07 decorative state"); body(d, "Informative, decorative, and informative images occur in that order."); alt_image(d,"fixture.png","Informative first image.","First"); alt_image(d,"decorative.png","","Decorative middle"); alt_image(d,"fixture.png","Informative last image.","Last"); save("B07", "decorative_image", d, {"figures":[{"alt":"Informative first image.","decorative":False,"order":1},{"alt":"","decorative":True,"order":2},{"alt":"Informative last image.","decorative":False,"order":3}]})
    d = setup("Variant B08 hyperlinks"); title(d,"Variant B08 hyperlinks"); p=d.add_paragraph(); p.add_run("Read "); hyperlink(p,"the WCAG standard","https://www.w3.org/TR/WCAG22/"); p.add_run(" and "); hyperlink(p,"the PDF reference","https://pdfa.org/"); save("B08", "links", d, {"hyperlinks":[{"text":"the WCAG standard","target":"https://www.w3.org/TR/WCAG22/","order":1},{"text":"the PDF reference","target":"https://pdfa.org/","order":2}]})
    d = setup("Variant B09: title / metadata"); d.core_properties.subject="Variant B09 subject: preservation"; d.core_properties.keywords="a11y, PDF, DOCX, 2026"; title(d,"Variant B09: title / metadata"); body(d,"Visible and core metadata values are declared separately."); save("B09", "document_title", d, {"document_identity":{"visible_title":"Variant B09: title / metadata","core_title":"Variant B09: title / metadata","subject":"Variant B09 subject: preservation","keywords":"a11y, PDF, DOCX, 2026"}})
    d = setup("Variant B10 complex table"); title(d,"Variant B10 complex table"); t=d.add_table(rows=5, cols=4); t.style="Table Grid"; t.cell(0,0).merge(t.cell(0,1)).text="Population"; t.cell(0,2).merge(t.cell(0,3)).text="Outcome"; t.cell(1,0).text="Year"; t.cell(1,1).text="Group"; t.cell(1,2).text="Rate"; t.cell(1,3).text="Count"; t.cell(2,0).text="2024"; t.cell(2,1).text="A"; t.cell(2,2).text="82%"; t.cell(2,3).text="12"; t.cell(3,0).text="2025"; t.cell(3,1).text="B"; t.cell(3,2).text="87%"; t.cell(3,3).text="18"; t.cell(4,0).text="2026"; t.cell(4,1).text="C"; t.cell(4,2).text="91%"; t.cell(4,3).text="22"; header_row(t,0); header_row(t,1); save("B10", "complex_table", d, {"table":{"dimensions":{"rows":5,"columns":4},"header_rows":[0,1],"cells":[["Population","Outcome"],["Year","Group","Rate","Count"],["2024","A","82%","12"],["2025","B","87%","18"],["2026","C","91%","22"]]}})
    d = setup("Variant B11 footnotes"); title(d,"Variant B11 footnotes"); body(d,"The first claim carries note one[[N1]], and the second carries note two[[N2]]."); temp=BASE/"_B11_temp.docx"; temp2=BASE/"_B11_temp2.docx"; out=FIXTURES/"B11.docx"; d.save(temp); insert_note(str(temp),str(temp2),"footnote","[[N1]]","First variant note."); insert_note(str(temp2),str(out),"footnote","[[N2]]","Second variant note."); temp.unlink(missing_ok=True); temp2.unlink(missing_ok=True); save_manifest = {"notes":[{"kind":"footnote","id":1,"text":"First variant note."},{"kind":"footnote","id":2,"text":"Second variant note."}]}; mp=MANIFESTS/"B11.json"; mp.write_text(json.dumps({"variant_id":"B11","feature_family":"footnotes","source_format":"DOCX","contract_version":"v3-variant-set-b","fixture_path":"corpus/v3_variant_set_b/fixtures/B11.docx","ground_truth":save_manifest},indent=2)+"\n",encoding="utf-8"); rows.append({"variant_id":"B11","feature_family":"footnotes","fixture":"corpus/v3_variant_set_b/fixtures/B11.docx","source_sha256":sha(out),"manifest":"corpus/v3_variant_set_b/manifests/B11.json","manifest_sha256":sha(mp)})
    d = setup("Variant B12 equations"); title(d,"Variant B12 equations"); body(d,"Two native math expressions are present."); p=d.add_paragraph(); math_expr(p,["E","=","m","c","2"]); p=d.add_paragraph(); math_expr(p,["a","=","b","+","c"]); save("B12", "equation", d, {"equations":[{"expression":"E = mc^2","tokens":["E","=","m","c","2"]},{"expression":"a = b + c","tokens":["a","=","b","+","c"]}]})
    d = setup("Variant B13 captions"); title(d,"Variant B13 captions"); body(d,"Two figure-caption associations are declared in source order."); alt_image(d,"fixture.png","First figure with workflow.","Figure one"); p=d.add_paragraph(style="Caption"); p.add_run("Figure 1. Workflow."); alt_image(d,"decorative.png","Second figure with results.","Figure two"); p=d.add_paragraph(style="Caption"); p.add_run("Figure 2. Results."); save("B13", "captions", d, {"figure_caption_pairs":[{"figure":"figure-1","caption":"Figure 1. Workflow.","order":1},{"figure":"figure-2","caption":"Figure 2. Results.","order":2}]})
    return rows

def sha(path: Path) -> str:
    h=hashlib.sha256();
    with path.open("rb") as f:
        for b in iter(lambda:f.read(1024*1024),b""): h.update(b)
    return h.hexdigest()

if __name__ == "__main__":
    FIXTURES.mkdir(parents=True, exist_ok=True); MANIFESTS.mkdir(parents=True, exist_ok=True)
    records=build(); (BASE/"BUILD_RECORD.json").write_text(json.dumps({"variant_set":"B","frozen":False,"records":records},indent=2)+"\n",encoding="utf-8"); print(json.dumps({"created":len(records),"records":records},indent=2))
