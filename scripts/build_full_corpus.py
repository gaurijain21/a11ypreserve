from __future__ import annotations

import hashlib
import json
import os
import shutil
import sys
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / "corpus" / "fixtures"
MANIFESTS = ROOT / "corpus" / "manifests"
ASSETS = ROOT / "corpus"
FOOTNOTE_HELPER = os.environ.get("A11YPRESERVE_DOCUMENT_SKILLS")
if FOOTNOTE_HELPER and FOOTNOTE_HELPER not in sys.path:
    sys.path.insert(0, FOOTNOTE_HELPER)
from insert_note import insert_note  # noqa: E402


W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"


def q(tag: str) -> str:
    return qn(tag)


def set_run_lang(run, lang: str) -> None:
    rpr = run._r.get_or_add_rPr()
    node = rpr.find(q("w:lang"))
    if node is None:
        node = OxmlElement("w:lang")
        rpr.append(node)
    node.set(q("w:val"), lang)


def set_document_defaults(document: Document, language: str = "en-US") -> None:
    styles = document.styles.element
    defaults = styles.find(q("w:docDefaults"))
    if defaults is None:
        defaults = OxmlElement("w:docDefaults")
        styles.insert(0, defaults)
    rpr_default = defaults.find(q("w:rPrDefault"))
    if rpr_default is None:
        rpr_default = OxmlElement("w:rPrDefault")
        defaults.append(rpr_default)
    rpr = rpr_default.find(q("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        rpr_default.append(rpr)
    lang = rpr.find(q("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(q("w:val"), language)
    lang.set(q("w:eastAsia"), language)
    lang.set(q("w:bidi"), language)


def setup_document(title: str, language: str = "en-US") -> Document:
    document = Document()
    document.core_properties.title = title
    document.core_properties.language = language
    document.core_properties.subject = "A11yPreserve Phase 4 controlled benchmark fixture"
    document.core_properties.author = "A11yPreserve benchmark"
    set_document_defaults(document, language)
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Caption"):
        style = document.styles[style_name]
        style.font.name = "Aptos"
        style.font.color.rgb = RGBColor(0, 0, 0)
        for border in list(style.element.xpath(".//w:pBdr")):
            border.getparent().remove(border)
    document.styles["Normal"].font.size = Pt(11)
    return document


def add_title(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="Title")
    run = paragraph.add_run(text)
    set_run_lang(run, "en-US")


def add_body(document: Document, text: str) -> Any:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    set_run_lang(run, "en-US")
    return paragraph


def add_heading(document: Document, text: str, level: int = 1) -> Any:
    paragraph = document.add_paragraph(style=f"Heading {level}")
    run = paragraph.add_run(text)
    set_run_lang(run, "en-US")
    return paragraph


def set_alt_text(inline, description: str, title: str = "") -> None:
    inline._inline.docPr.set("descr", description)
    inline._inline.docPr.set("title", title)


def add_image(document: Document, path: Path, alt: str, title: str = "") -> Any:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    inline = paragraph.add_run().add_picture(str(path), width=Inches(2.3))
    set_alt_text(inline, alt, title)
    return paragraph


def add_hyperlink(paragraph, text: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(q("r:id"), relationship_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(q("w:val"), "0000FF")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(q("w:val"), "single")
    rpr.append(underline)
    run.append(rpr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def mark_header_row(table, row_index: int) -> None:
    tr_pr = table.rows[row_index]._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(q("w:val"), "true")
    tr_pr.append(header)


def add_math_token(parent, text: str) -> None:
    run = OxmlElement("m:r")
    text_node = OxmlElement("m:t")
    text_node.text = text
    run.append(text_node)
    parent.append(run)


def add_equation(paragraph) -> None:
    para = OxmlElement("m:oMathPara")
    math = OxmlElement("m:oMath")
    add_math_token(math, "E")
    add_math_token(math, "=")
    add_math_token(math, "m")
    sup = OxmlElement("m:sSup")
    base = OxmlElement("m:e")
    add_math_token(base, "c")
    exponent = OxmlElement("m:sup")
    add_math_token(exponent, "2")
    sup.append(base)
    sup.append(exponent)
    math.append(sup)
    para.append(math)
    paragraph._p.append(para)


def build_f05() -> tuple[Document, dict[str, Any]]:
    title = "A11yPreserve Document Language Fixture"
    document = setup_document(title, "en-US")
    add_title(document, title)
    add_body(document, "This document declares English United States as its primary language.")
    return document, {"document_language": "en-US"}


def build_f06() -> tuple[Document, dict[str, Any]]:
    title = "A11yPreserve Inline Language Fixture"
    document = setup_document(title, "en-US")
    add_title(document, title)
    paragraph = document.add_paragraph()
    english = paragraph.add_run("The greeting is ")
    set_run_lang(english, "en-US")
    spanish = paragraph.add_run("Buenos días")
    set_run_lang(spanish, "es-MX")
    tail = paragraph.add_run(" before the meeting.")
    set_run_lang(tail, "en-US")
    return document, {"document_language": "en-US", "runs": [{"text": "Buenos días", "language": "es-MX"}]}


def build_f07() -> tuple[Document, dict[str, Any]]:
    title = "A11yPreserve Decorative Image Fixture"
    document = setup_document(title)
    add_title(document, title)
    add_body(document, "The first image conveys information; the second image is explicitly decorative.")
    meaningful = add_image(document, ASSETS / "fixture.png", "A meaningful research figure showing the benchmark workflow.", "Informative figure")
    decorative = add_image(document, ASSETS / "decorative.png", "", "Decorative image")
    return document, {
        "figures": [
            {"identifier": "figure-1", "alt": "A meaningful research figure showing the benchmark workflow.", "decorative": False, "order": 1},
            {"identifier": "figure-2", "alt": "", "decorative": True, "order": 2},
        ],
    }


def build_f08() -> tuple[Document, dict[str, Any]]:
    title = "A11yPreserve Hyperlink Fixture"
    document = setup_document(title)
    add_title(document, title)
    paragraph = document.add_paragraph()
    prefix = paragraph.add_run("Read the ")
    set_run_lang(prefix, "en-US")
    add_hyperlink(paragraph, "Accessibility Guidelines", "https://www.w3.org/WAI/standards-guidelines/wcag/")
    return document, {"hyperlinks": [{"text": "Accessibility Guidelines", "target": "https://www.w3.org/WAI/standards-guidelines/wcag/", "order": 1}]}


def build_f09() -> tuple[Document, dict[str, Any]]:
    title = "A11yPreserve Accessible Research Brief"
    document = setup_document(title)
    document.core_properties.subject = "Controlled document identity metadata"
    document.core_properties.keywords = "accessibility, conversion, preservation"
    add_title(document, title)
    add_body(document, "The visible title and core document title identify the same research brief.")
    return document, {
        "document_identity": {
            "visible_title": title,
            "core_title": title,
            "subject": "Controlled document identity metadata",
            "keywords": "accessibility, conversion, preservation",
        }
    }


def build_f10() -> tuple[Document, dict[str, Any]]:
    title = "A11yPreserve Complex Table Fixture"
    document = setup_document(title)
    add_title(document, title)
    add_body(document, "The first two rows are headers. The first row groups the two enrollment columns.")
    table = document.add_table(rows=5, cols=3)
    table.style = "Table Grid"
    top_left = table.cell(0, 0).merge(table.cell(0, 1))
    top_left.text = "Enrollment"
    table.cell(0, 2).text = "Outcome"
    headers = ["Year", "Students", "Completion Rate"]
    for index, value in enumerate(headers):
        table.cell(1, index).text = value
    for row_index, values in enumerate([["2024", "120", "82%"], ["2025", "150", "87%"], ["2026", "180", "91%"]], start=2):
        for col_index, value in enumerate(values):
            table.cell(row_index, col_index).text = value
    mark_header_row(table, 0)
    mark_header_row(table, 1)
    return document, {
        "table": {
            "dimensions": {"rows": 5, "columns": 3},
            "header_rows": [0, 1],
            "header_groups": [{"text": "Enrollment", "columns": [0, 1]}, {"text": "Outcome", "columns": [2]}],
            "headers": ["Year", "Students", "Completion Rate"],
            "cells": [["Enrollment", "Outcome"], headers, ["2024", "120", "82%"], ["2025", "150", "87%"], ["2026", "180", "91%"]],
        }
    }


def build_f11() -> tuple[Document, dict[str, Any]]:
    title = "A11yPreserve Footnote Fixture"
    document = setup_document(title)
    add_title(document, title)
    add_body(document, "The benchmark uses an explicit footnote association[[FN]].")
    temp = ROOT / "corpus" / "_F11_footnote_source.docx"
    output = FIXTURES / "F11_FOOTNOTES.docx"
    document.save(temp)
    insert_note(str(temp), str(output), "footnote", "[[FN]]", "This note records the source-to-note association.")
    temp.unlink(missing_ok=True)
    return Document(str(output)), {"notes": [{"kind": "footnote", "reference_text": "The benchmark uses an explicit footnote association.", "id": 1, "text": "This note records the source-to-note association."}]}


def build_f12() -> tuple[Document, dict[str, Any]]:
    title = "A11yPreserve Equation Fixture"
    document = setup_document(title)
    add_title(document, title)
    add_body(document, "The fixture contains a native OMML equation rather than an image substitute.")
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_equation(paragraph)
    return document, {"equations": [{"expression": "E = mc^2", "representation": "OMML", "tokens": ["E", "=", "m", "c", "2"]}]}


def build_f14() -> tuple[Document, dict[str, Any]]:
    title = "A11yPreserve Figure Caption Fixture"
    document = setup_document(title)
    add_title(document, title)
    add_body(document, "The figure and its caption form one source-level association.")
    add_image(document, ASSETS / "fixture.png", "A meaningful figure showing the benchmark workflow.", "Benchmark figure")
    caption = document.add_paragraph(style="Caption")
    run = caption.add_run("Figure 1. Benchmark workflow overview.")
    set_run_lang(run, "en-US")
    return document, {"figure_caption_pairs": [{"figure": "figure-1", "caption": "Figure 1. Benchmark workflow overview.", "order": 1}]}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def main() -> None:
    FIXTURES.mkdir(parents=True, exist_ok=True)
    MANIFESTS.mkdir(parents=True, exist_ok=True)
    builders = [
        ("F05_DOCUMENT_LANGUAGE", "A11yPreserve Document Language Fixture", build_f05),
        ("F06_INLINE_LANGUAGE", "A11yPreserve Inline Language Fixture", build_f06),
        ("F07_DECORATIVE_IMAGE", "A11yPreserve Decorative Image Fixture", build_f07),
        ("F08_LINKS", "A11yPreserve Hyperlink Fixture", build_f08),
        ("F09_DOCUMENT_TITLE", "A11yPreserve Accessible Research Brief", build_f09),
        ("F10_COMPLEX_TABLE", "A11yPreserve Complex Table Fixture", build_f10),
        ("F11_FOOTNOTES", "A11yPreserve Footnote Fixture", build_f11),
        ("F12_EQUATION", "A11yPreserve Equation Fixture", build_f12),
        ("F14_CAPTIONS", "A11yPreserve Figure Caption Fixture", build_f14),
    ]
    records: list[dict[str, Any]] = []
    for fixture_id, title, builder in builders:
        output = FIXTURES / f"{fixture_id}.docx"
        if fixture_id != "F11_FOOTNOTES":
            document, ground_truth = builder()
            document.save(output)
        else:
            document, ground_truth = builder()
        manifest = {"fixture_id": fixture_id, "feature": fixture_id.removeprefix("F05_").removeprefix("F06_").removeprefix("F07_").removeprefix("F08_").removeprefix("F09_").removeprefix("F10_").removeprefix("F11_").removeprefix("F12_").removeprefix("F14_").lower(), "source_format": "DOCX", "contract_version": "phase4-v1", "document": {"title": title, "language": "en-US"}, "ground_truth": ground_truth}
        manifest_path = MANIFESTS / f"{fixture_id}.json"
        manifest_path.write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
        records.append({"fixture": fixture_id, "path": str(output), "size_bytes": output.stat().st_size, "sha256": sha256(output), "manifest": str(manifest_path), "manifest_sha256": sha256(manifest_path), "verification_status": "PENDING_OOXML_VERIFICATION"})
    (ROOT / "corpus" / "BUILD_RECORD.json").write_text(json.dumps({"generated": records, "frozen": False}, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"created": records}, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
