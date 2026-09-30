from __future__ import annotations

import hashlib
import json
import os
import sys
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "integrated_study"
DOCX_DIR = OUT / "sources"
ASSETS = ROOT / "corpus"
FOOTNOTE_HELPER = os.environ.get("A11YPRESERVE_DOCUMENT_SKILLS")
if FOOTNOTE_HELPER and FOOTNOTE_HELPER not in sys.path:
    sys.path.insert(0, FOOTNOTE_HELPER)
from insert_note import insert_note  # noqa: E402


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def q(tag: str) -> str:
    return qn(tag)


def set_run_language(run, language: str) -> None:
    rpr = run._r.get_or_add_rPr()
    node = rpr.find(q("w:lang"))
    if node is None:
        node = OxmlElement("w:lang")
        rpr.append(node)
    node.set(q("w:val"), language)


def set_document_language(document: Document, language: str) -> None:
    styles = document.styles.element
    defaults = styles.find(q("w:docDefaults"))
    if defaults is None:
        defaults = OxmlElement("w:docDefaults")
        styles.insert(0, defaults)
    rpr_default = defaults.find(q("w:rPrDefault"))
    if rpr_default is None:
        rpr_default = OxmlElement("w:rPrDefault")
        defaults.append(rpr_default)
    rpr = rpr_default.find(q("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        rpr_default.append(rpr)
    node = rpr.find(q("w:lang"))
    if node is None:
        node = OxmlElement("w:lang")
        rpr.append(node)
    node.set(q("w:val"), language)
    node.set(q("w:eastAsia"), language)
    node.set(q("w:bidi"), language)


def setup_document(title: str, language: str = "en-US") -> Document:
    document = Document()
    document.core_properties.title = title
    document.core_properties.language = language
    document.core_properties.subject = "A11yPreserve integrated external-validity sanity check"
    document.core_properties.author = "A11yPreserve benchmark"
    set_document_language(document, language)
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Caption"):
        style = document.styles[style_name]
        style.font.name = "Aptos"
        style.font.color.rgb = RGBColor(0, 0, 0)
    document.styles["Normal"].font.size = Pt(10.5)
    return document


def add_title(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="Title")
    run = paragraph.add_run(text)
    set_run_language(run, "en-US")


def add_heading(document: Document, text: str, level: int = 1):
    paragraph = document.add_paragraph(style=f"Heading {level}")
    run = paragraph.add_run(text)
    set_run_language(run, "en-US")
    return paragraph


def add_body(document: Document, text: str):
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    set_run_language(run, "en-US")
    return paragraph


def add_language_body(document: Document, before: str, foreign: str, after: str, language: str):
    paragraph = document.add_paragraph()
    first = paragraph.add_run(before)
    set_run_language(first, "en-US")
    second = paragraph.add_run(foreign)
    set_run_language(second, language)
    third = paragraph.add_run(after)
    set_run_language(third, "en-US")
    return paragraph


def add_list(document: Document, items: list[tuple[str, int]], style: str = "List Bullet") -> None:
    for text, level in items:
        paragraph = document.add_paragraph(style=style)
        paragraph.paragraph_format.left_indent = Inches(0.25 * level)
        run = paragraph.add_run(text)
        set_run_language(run, "en-US")
        ppr = paragraph._p.get_or_add_pPr()
        num_pr = ppr.find(q("w:numPr"))
        if num_pr is None:
            num_pr = OxmlElement("w:numPr")
            ppr.append(num_pr)
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(q("w:val"), str(level))
        num_id = OxmlElement("w:numId")
        num_id.set(q("w:val"), "1")
        num_pr.append(ilvl)
        num_pr.append(num_id)


def mark_header_row(table, row_index: int) -> None:
    tr_pr = table.rows[row_index]._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(q("w:val"), "true")
    tr_pr.append(header)


def add_table(document: Document, headers: list[str], rows: list[list[str]], title: str):
    add_body(document, title)
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    mark_header_row(table, 0)
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
    return table


def add_hyperlink(paragraph, text: str, target: str) -> None:
    relationship_id = paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(q("r:id"), relationship_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(q("w:val"), "0000FF")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(q("w:val"), "single")
    rpr.append(underline)
    run.append(rpr)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_link_paragraph(document: Document, prefix: str, text: str, target: str):
    paragraph = document.add_paragraph()
    run = paragraph.add_run(prefix)
    set_run_language(run, "en-US")
    add_hyperlink(paragraph, text, target)
    return paragraph


def add_image(document: Document, alt: str, title: str, decorative: bool = False):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image = ASSETS / ("decorative.png" if decorative else "fixture.png")
    inline = paragraph.add_run().add_picture(str(image), width=Inches(2.4))
    inline._inline.docPr.set("descr", alt)
    inline._inline.docPr.set("title", title)
    return paragraph


def add_equation(paragraph) -> None:
    para = OxmlElement("m:oMathPara")
    math = OxmlElement("m:oMath")
    for token in ("E", "=", "m"):
        run = OxmlElement("m:r")
        node = OxmlElement("m:t")
        node.text = token
        run.append(node)
        math.append(run)
    sup = OxmlElement("m:sSup")
    base = OxmlElement("m:e")
    run = OxmlElement("m:r")
    node = OxmlElement("m:t")
    node.text = "c"
    run.append(node)
    base.append(run)
    exponent = OxmlElement("m:sup")
    run = OxmlElement("m:r")
    node = OxmlElement("m:t")
    node.text = "2"
    run.append(node)
    exponent.append(run)
    sup.append(base)
    sup.append(exponent)
    math.append(sup)
    para.append(math)
    paragraph._p.append(para)


def add_caption(document: Document, text: str):
    paragraph = document.add_paragraph(style="Caption")
    run = paragraph.add_run(text)
    set_run_language(run, "en-US")
    return paragraph


def prose(document: Document, paragraphs: list[str]) -> None:
    for text in paragraphs:
        add_body(document, text)


def common_report_prose(document: Document) -> None:
    prose(document, [
        "This report describes a small accessibility-preservation review for a public-service program. It records the intended structure before publication so that later format conversion can be checked against the authoring contract.",
        "The review team combines narrative explanation with structured findings. Headings divide the report into purpose, evidence, findings, and next steps; the table gives readers a compact way to compare the same measures across reporting periods.",
        "The document is designed for readers who may navigate by headings, lists, tables, links, and image alternatives. Those features are present as authored document information rather than being inferred only from visual appearance.",
    ])


def build_i01() -> tuple[Document, list[dict[str, Any]]]:
    title = "Community Access Report: Service Improvements and Next Steps"
    document = setup_document(title)
    add_title(document, title)
    add_body(document, "Prepared for the Community Access Office | Reporting period: 2025")
    add_heading(document, "Executive summary", 1)
    common_report_prose(document)
    add_heading(document, "What changed this year", 1)
    add_heading(document, "Service access", 2)
    prose(document, [
        "The office expanded appointment hours, simplified intake language, and published a single contact route for questions. Staff reviewed the changes with community partners and recorded follow-up actions in the service log.",
        "The most useful indicator is not the number of redesigned pages by itself. It is whether people can locate the right service, understand the next step, and use the information in the format that was delivered.",
    ])
    add_heading(document, "Community feedback", 2)
    add_list(document, [
        ("Keep the appointment checklist close to the service description.", 0),
        ("Explain required documents before the first visit.", 0),
        ("Offer a direct route for language and format requests.", 1),
        ("Publish contact hours in the same place as the telephone number.", 1),
        ("Review the checklist after each quarterly update.", 0),
    ])
    add_heading(document, "Measures", 1)
    add_table(document, ["Measure", "2024", "2025"], [["Completed appointments", "1,240", "1,486"], ["Average wait days", "8", "5"], ["Follow-up requests", "164", "139"]], "Table 1. Service access measures.")
    add_heading(document, "Visual evidence", 1)
    add_body(document, "The figure summarizes the direction of the reporting trend and is accompanied by an author-written alternative description.")
    add_image(document, "A simple upward trend showing completed appointments increasing from 2024 to 2025.", "Appointments trend")
    add_caption(document, "Figure 1. Completed appointments increased during the reporting period.")
    add_link_paragraph(document, "The public service guidance is available in the ", "community access handbook", "https://www.w3.org/WAI/standards-guidelines/wcag/")
    add_heading(document, "Next steps", 1)
    prose(document, [
        "The next reporting cycle will retain the same headings and measure definitions so that changes can be compared without losing the context of the earlier report.",
        "Managers should confirm that the published document, its downloadable copy, and any converted PDF retain the information that was intentionally supplied for accessibility.",
    ])
    contracts = [
        {"instance": "I01-heading-hierarchy", "feature": "headings", "expected": {"levels": [1, 1, 2, 2, 1, 1, 1]}},
        {"instance": "I01-alternative-text", "feature": "image_alt_text", "expected": {"alt": "A simple upward trend showing completed appointments increasing from 2024 to 2025."}},
        {"instance": "I01-lists", "feature": "lists", "expected": {"depths": [0, 0, 1, 1, 0]}},
        {"instance": "I01-table-headers", "feature": "table", "expected": {"headers": ["Measure", "2024", "2025"], "rows": 4}},
        {"instance": "I01-document-language", "feature": "document_language", "expected": {"language": "en-US"}},
        {"instance": "I01-hyperlink", "feature": "hyperlinks", "expected": {"target": "https://www.w3.org/WAI/standards-guidelines/wcag/"}},
        {"instance": "I01-title", "feature": "document_title", "expected": {"title": title}},
    ]
    return document, contracts


def build_i02() -> tuple[Document, list[dict[str, Any]]]:
    title = "Accessible Learning Handout: Reading Evidence in Public Reports"
    document = setup_document(title)
    add_title(document, title)
    add_body(document, "An academic handout for a short course on evidence, interpretation, and inclusive communication.")
    add_heading(document, "Learning goals", 1)
    prose(document, [
        "By the end of the session, learners should be able to distinguish an observation from an interpretation, identify the structure that carries an author’s meaning, and explain why a converted document may require a second inspection.",
        "The handout uses a short bilingual example because language changes are part of document meaning. The foreign-language phrase is intentionally marked in the source document rather than treated as decoration.",
    ])
    add_language_body(document, "The example greeting is ", "Buenos días", " before the discussion begins.", "es-MX")
    add_heading(document, "Reading sequence", 1)
    add_list(document, [("Identify the claim.", 0), ("Locate the evidence.", 0), ("Check the relationship between evidence and conclusion.", 1), ("Record uncertainty instead of filling it with assumption.", 1), ("Explain the result to another reader.", 0)], style="List Number")
    add_heading(document, "A compact model", 1)
    add_body(document, "The following expression is included as a native equation so that the mathematical relationship is part of the document structure.")
    equation = document.add_paragraph()
    equation.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_equation(equation)
    add_caption(document, "Equation 1. A compact expression used in the classroom example.")
    add_heading(document, "Case discussion", 1)
    prose(document, [
        "A reader may find the same idea in a paragraph, a table, a figure, or a note. The representation changes, but the contract is still about the relationship that the author deliberately supplied.",
        "When a document is converted, visible text alone may survive while the structure that identifies a relationship does not. The class therefore compares both content and machine-verifiable structure.",
    ])
    add_body(document, "The handout uses an explicit footnote to identify the source of the classroom example[[FN]].")
    temp = OUT / "_I02_footnote_source.docx"
    raw = OUT / "sources" / "I02_EDUCATIONAL_HANDOUT_raw.docx"
    document.save(temp)
    insert_note(str(temp), str(raw), "footnote", "[[FN]]", "This note identifies the source used for the classroom example.")
    temp.unlink(missing_ok=True)
    document = Document(str(raw))
    add_heading(document, "Figure and caption", 1)
    add_body(document, "The figure gives the class a visual summary; the caption states what the figure is intended to communicate.")
    add_image(document, "A classroom diagram showing evidence connected to a conclusion.", "Evidence diagram")
    add_caption(document, "Figure 1. Evidence and conclusion are connected through an explicit reasoning step.")
    add_link_paragraph(document, "Further reading is listed in the ", "accessibility standards overview", "https://www.w3.org/WAI/standards-guidelines/wcag/")
    prose(document, [
        "The final exercise asks students to describe what is known, what is not known, and what additional evidence would be needed. That discipline is part of accessible communication because it prevents readers from mistaking an uncertain representation for a confirmed one.",
        "The handout ends with a short reflection: Which parts of the source meaning are carried by words, and which parts are carried by structure?",
    ])
    contracts = [
        {"instance": "I02-heading-hierarchy", "feature": "headings", "expected": {"levels": [1, 1, 1, 1, 1]}},
        {"instance": "I02-inline-language", "feature": "inline_language", "expected": {"text": "Buenos días", "language": "es-MX"}},
        {"instance": "I02-equation", "feature": "equation", "expected": {"omml": True}},
        {"instance": "I02-footnote", "feature": "footnotes", "expected": {"text": "This note identifies the source used for the classroom example."}},
        {"instance": "I02-caption", "feature": "captions", "expected": {"text": "Figure 1. Evidence and conclusion are connected through an explicit reasoning step."}},
        {"instance": "I02-hyperlink", "feature": "hyperlinks", "expected": {"target": "https://www.w3.org/WAI/standards-guidelines/wcag/"}},
    ]
    return document, contracts


def build_i03() -> tuple[Document, list[dict[str, Any]]]:
    title = "Plain-Language Policy Note: Updating Public Information"
    document = setup_document(title)
    add_title(document, title)
    add_body(document, "Policy and implementation note | Version 2.0 | 2025")
    add_heading(document, "Purpose", 1)
    prose(document, [
        "This note explains how an agency should update public information while preserving the structure that helps readers find, understand, and reuse it. The policy applies to web pages, downloadable documents, and converted PDF copies.",
        "The note is intentionally written in plain language. Short sections, descriptive headings, lists, and tables make the decision path visible without requiring a reader to infer it from layout alone.",
    ])
    add_heading(document, "Policy requirements", 1)
    add_heading(document, "Before publication", 2)
    add_list(document, [("Identify the audience and the decision the document supports.", 0), ("Write a short title and a meaningful summary.", 0), ("Check every image for an author-supplied alternative or a documented decorative state.", 1), ("Check table headers and links before release.", 1)], style="List Bullet")
    add_heading(document, "During review", 2)
    add_list(document, [("Review the heading hierarchy.", 0), ("Confirm the primary document language.", 0), ("Inspect converted outputs against the source contract.", 0)], style="List Bullet")
    add_heading(document, "Implementation schedule", 1)
    add_table(document, ["Phase", "Owner", "Evidence"], [["Draft", "Policy team", "Source contract"], ["Review", "Accessibility lead", "Issue log"], ["Release", "Publishing team", "Converted copy"]], "Table 1. Policy implementation schedule.")
    add_heading(document, "Illustrations", 1)
    add_body(document, "The first illustration carries information; the second is decorative and is included only to separate content from presentation.")
    add_image(document, "A process diagram showing draft, review, and release stages.", "Policy process diagram")
    add_caption(document, "Figure 1. The policy moves through draft, review, and release stages.")
    add_image(document, "", "Decorative divider", decorative=True)
    add_heading(document, "Language and links", 1)
    add_language_body(document, "The policy uses the English document language and identifies the phrase ", "información pública", " as Spanish when it appears in an example.", "es-MX")
    add_link_paragraph(document, "The release checklist is aligned with the ", "Web Content Accessibility Guidelines", "https://www.w3.org/WAI/standards-guidelines/wcag/")
    add_heading(document, "Review record", 1)
    prose(document, [
        "Reviewers should record both positive and negative findings. A preserved heading does not prove that a table association survived, and a tagged PDF does not by itself prove that every known source property survived the conversion.",
        "The release record should retain the source hash, conversion conditions, destination hash, and evidence path for each checked property. This makes later corrections possible without rewriting the original decision.",
    ])
    contracts = [
        {"instance": "I03-heading-hierarchy", "feature": "headings", "expected": {"levels": [1, 1, 2, 2, 1, 1, 1, 1]}},
        {"instance": "I03-lists", "feature": "lists", "expected": {"depths": [0, 0, 1, 1, 0, 0, 0]}},
        {"instance": "I03-table-headers", "feature": "table", "expected": {"headers": ["Phase", "Owner", "Evidence"], "rows": 4}},
        {"instance": "I03-image-alt", "feature": "image_alt_text", "expected": {"alt": "A process diagram showing draft, review, and release stages."}},
        {"instance": "I03-decorative-image", "feature": "decorative_image", "expected": {"decorative": True}},
        {"instance": "I03-document-language", "feature": "document_language", "expected": {"language": "en-US"}},
        {"instance": "I03-inline-language", "feature": "inline_language", "expected": {"text": "información pública", "language": "es-MX"}},
        {"instance": "I03-hyperlink", "feature": "hyperlinks", "expected": {"target": "https://www.w3.org/WAI/standards-guidelines/wcag/"}},
    ]
    return document, contracts


def main() -> None:
    DOCX_DIR.mkdir(parents=True, exist_ok=True)
    builders = [("I01_REALISTIC_REPORT", build_i01), ("I02_EDUCATIONAL_HANDOUT", build_i02), ("I03_POLICY_NOTE", build_i03)]
    records = []
    for document_id, builder in builders:
        document, contracts = builder()
        path = DOCX_DIR / f"{document_id}.docx"
        if not path.exists():
            document.save(path)
        else:
            document.save(path)
        records.append({"document_id": document_id, "source_docx": str(path.relative_to(ROOT)), "source_sha256": sha256(path), "contracts": contracts})
    (OUT / "INTEGRATED_CONTRACTS.json").write_text(json.dumps({"documents": records}, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps({"documents": len(records), "contracts": sum(len(r["contracts"]) for r in records)}, indent=2))


if __name__ == "__main__":
    main()
