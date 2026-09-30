from __future__ import annotations

import hashlib
import json
import os
import platform
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

try:
    from PIL import Image, ImageDraw, ImageFont
except ImportError as exc:  # pragma: no cover
    raise SystemExit("Pillow is required to create the local chart fixture") from exc


ROOT = Path(__file__).resolve().parents[1]
PILOT = ROOT / "pilot"
FIXTURES = PILOT / "fixtures"
MANIFESTS = PILOT / "manifests"
LOGS = PILOT / "logs"


def q(tag: str) -> str:
    return qn(tag)


def set_run_lang(run, lang: str = "en-US") -> None:
    rpr = run._r.get_or_add_rPr()
    node = rpr.find(q("w:lang"))
    if node is None:
        node = OxmlElement("w:lang")
        rpr.append(node)
    node.set(q("w:val"), lang)


def set_paragraph_lang(paragraph, lang: str = "en-US") -> None:
    for run in paragraph.runs:
        set_run_lang(run, lang)


def set_document_defaults(document: Document) -> None:
    styles = document.styles.element
    defaults = styles.find(q("w:docDefaults"))
    if defaults is None:
        defaults = OxmlElement("w:docDefaults")
        styles.insert(0, defaults)
    rpr_default = defaults.find(q("w:rPrDefault"))
    if rpr_default is None:
        rpr_default = OxmlElement("w:rPrDefault")
        defaults.append(rpr_default)
    rpr = rpr_default.find(q("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        rpr_default.append(rpr)
    lang = rpr.find(q("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(q("w:val"), "en-US")
    lang.set(q("w:eastAsia"), "en-US")
    lang.set(q("w:bidi"), "en-US")


def set_style_black(style) -> None:
    style.font.name = "Aptos"
    style.font.color.rgb = RGBColor(0, 0, 0)


def setup_document(title: str) -> Document:
    document = Document()
    document.core_properties.title = title
    document.core_properties.language = "en-US"
    document.core_properties.subject = "A11yPreserve controlled accessibility fixture"
    document.core_properties.author = "A11yPreserve pilot"
    set_document_defaults(document)
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        set_style_black(document.styles[style_name])
    document.styles["Normal"].font.size = Pt(11)
    return document


def add_title(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="Title")
    run = paragraph.add_run(text)
    set_run_lang(run)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT


def add_body(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(text)
    set_paragraph_lang(paragraph)


def add_heading(document: Document, text: str, level: int) -> None:
    paragraph = document.add_paragraph(text, style=f"Heading {level}")
    set_paragraph_lang(paragraph)


def create_numbering(document: Document, fmt: str, text: str) -> int:
    numbering = document.part.numbering_part.element
    abstract_ids = [int(x.get(q("w:abstractNumId"))) for x in numbering.findall(q("w:abstractNum"))]
    num_ids = [int(x.get(q("w:numId"))) for x in numbering.findall(q("w:num"))]
    abstract_id = max(abstract_ids or [0]) + 1
    num_id = max(num_ids or [0]) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(q("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(q("w:val"), "multilevel")
    abstract.append(multi)
    for level in range(3):
        lvl = OxmlElement("w:lvl")
        lvl.set(q("w:ilvl"), str(level))
        start = OxmlElement("w:start")
        start.set(q("w:val"), "1")
        lvl.append(start)
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(q("w:val"), fmt)
        lvl.append(num_fmt)
        lvl_text = OxmlElement("w:lvlText")
        if fmt == "bullet":
            lvl_text.set(q("w:val"), "•")
        else:
            lvl_text.set(q("w:val"), f"%{level + 1}.")
        lvl.append(lvl_text)
        suff = OxmlElement("w:suff")
        suff.set(q("w:val"), "tab")
        lvl.append(suff)
        ppr = OxmlElement("w:pPr")
        ind = OxmlElement("w:ind")
        ind.set(q("w:left"), str(720 + 360 * level))
        ind.set(q("w:hanging"), "360")
        ppr.append(ind)
        lvl.append(ppr)
        rpr = OxmlElement("w:rPr")
        if fmt == "bullet":
            rfonts = OxmlElement("w:rFonts")
            rfonts.set(q("w:ascii"), "Symbol")
            rfonts.set(q("w:hAnsi"), "Symbol")
            rpr.append(rfonts)
        lvl.append(rpr)
        abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(q("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(q("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def set_list_paragraph(paragraph, num_id: int, level: int) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = ppr.find(q("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        ppr.append(num_pr)
    ilvl = num_pr.find(q("w:ilvl"))
    if ilvl is None:
        ilvl = OxmlElement("w:ilvl")
        num_pr.append(ilvl)
    ilvl.set(q("w:val"), str(level))
    num = num_pr.find(q("w:numId"))
    if num is None:
        num = OxmlElement("w:numId")
        num_pr.append(num)
    num.set(q("w:val"), str(num_id))
    set_paragraph_lang(paragraph)


def set_alt_text(inline, description: str) -> None:
    inline._inline.docPr.set("descr", description)
    inline._inline.docPr.set("title", "Enrollment chart")


def mark_header_row(table) -> None:
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(q("w:val"), "true")
    tr_pr.append(header)
    for cell in table.rows[0].cells:
        for p in cell.paragraphs:
            set_paragraph_lang(p)


def make_chart(path: Path) -> None:
    image = Image.new("RGB", (1200, 700), "white")
    draw = ImageDraw.Draw(image)
    margin_left, margin_bottom, top = 130, 110, 70
    chart_bottom, chart_top = 590, top
    chart_right = 1110
    draw.text((55, 20), "Enrollment trend", fill="black")
    draw.line((margin_left, chart_top, margin_left, chart_bottom), fill="black", width=3)
    draw.line((margin_left, chart_bottom, chart_right, chart_bottom), fill="black", width=3)
    values = [("2024", 120), ("2025", 150), ("2026", 180)]
    max_value = 200
    bar_width = 180
    gap = 90
    for i, (year, value) in enumerate(values):
        x = margin_left + 90 + i * (bar_width + gap)
        y = chart_bottom - int((chart_bottom - chart_top) * value / max_value)
        draw.rectangle((x, y, x + bar_width, chart_bottom), fill=(55, 115, 180), outline="black", width=2)
        draw.text((x + 65, chart_bottom + 25), year, fill="black")
        draw.text((x + 60, y - 35), str(value), fill="black")
    draw.text((35, 250), "Students", fill="black")
    image.save(path, format="PNG")


def make_headings() -> tuple[Document, dict[str, Any]]:
    title = "A11yPreserve Heading Fixture"
    d = setup_document(title)
    add_title(d, title)
    for level, heading, body in [
        (1, "Research Overview", "This fixture introduces the research problem and establishes a deliberate heading hierarchy."),
        (2, "Background", "Accessible structure can be preserved only when conversion retains the source document's intended semantics."),
        (3, "Prior Work", "Prior accessibility research shows that document transformation can change structure and metadata."),
        (2, "Methods", "The pilot compares source semantics with tagged PDF semantics using deterministic extraction."),
    ]:
        add_heading(d, heading, level)
        add_body(d, body)
    return d, {"headings": [{"text": h, "level": l} for l, h, _ in [(1, "Research Overview", ""), (2, "Background", ""), (3, "Prior Work", ""), (2, "Methods", "")]]}


def make_alt_text(chart_path: Path) -> tuple[Document, dict[str, Any]]:
    title = "A11yPreserve Alternative Text Fixture"
    alt = "Bar chart showing enrollment increasing from 120 students in 2024 to 180 students in 2026."
    d = setup_document(title)
    add_title(d, title)
    add_body(d, "Enrollment is reported across three academic years in the chart below.")
    inline = d.add_picture(str(chart_path), width=Inches(5.8))
    set_alt_text(inline, alt)
    p = d.paragraphs[-1]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_body(d, "The visual summarizes the reported trend without repeating its complete alternative text.")
    return d, {"figures": [{"identifier": "figure-1", "alt": alt, "decorative": False, "order": 1}]}


def make_lists() -> tuple[Document, dict[str, Any]]:
    title = "A11yPreserve List Fixture"
    d = setup_document(title)
    add_title(d, title)
    add_body(d, "The fixture contains separate unordered, ordered, and nested lists with native Word numbering semantics.")
    bullet_id = create_numbering(d, "bullet", "•")
    number_id = create_numbering(d, "decimal", "%1.")
    p = d.add_paragraph("Research areas")
    set_paragraph_lang(p)
    rows = [("Accessibility", 0), ("Documents", 1), ("Web", 1), ("Usability", 0)]
    for text, level in rows:
        p = d.add_paragraph(text, style="List Bullet")
        set_list_paragraph(p, bullet_id, level)
    p = d.add_paragraph("Evaluation stages")
    set_paragraph_lang(p)
    stages = ["Extract", "Convert", "Compare"]
    for text in stages:
        p = d.add_paragraph(text, style="List Number")
        set_list_paragraph(p, number_id, 0)
    return d, {
        "lists": [
            {"type": "unordered", "depth": 0, "text": "Accessibility", "parent": None},
            {"type": "unordered", "depth": 1, "text": "Documents", "parent": "Accessibility"},
            {"type": "unordered", "depth": 1, "text": "Web", "parent": "Accessibility"},
            {"type": "unordered", "depth": 0, "text": "Usability", "parent": None},
            {"type": "ordered", "depth": 0, "text": "Extract", "parent": None},
            {"type": "ordered", "depth": 0, "text": "Convert", "parent": None},
            {"type": "ordered", "depth": 0, "text": "Compare", "parent": None},
        ]
    }


def make_table() -> tuple[Document, dict[str, Any]]:
    title = "A11yPreserve Table Fixture"
    d = setup_document(title)
    add_title(d, title)
    add_body(d, "The table reports enrollment and completion rate with the first row marked as a native Word header row.")
    table = d.add_table(rows=4, cols=3)
    table.style = "Table Grid"
    values = [
        ["Year", "Students", "Completion Rate"],
        ["2024", "120", "82%"],
        ["2025", "150", "87%"],
        ["2026", "180", "91%"],
    ]
    for row, vals in zip(table.rows, values):
        for cell, value in zip(row.cells, vals):
            cell.text = value
            for p in cell.paragraphs:
                set_paragraph_lang(p)
    mark_header_row(table)
    return d, {
        "table": {
            "dimensions": {"rows": 4, "columns": 3},
            "headers": ["Year", "Students", "Completion Rate"],
            "header_row": 0,
            "cells": values,
        }
    }


def manifest(fixture_id: str, title: str, features: dict[str, Any]) -> dict[str, Any]:
    result = {"fixture_id": fixture_id, "document": {"title": title, "language": "en-US"}}
    result.update(features)
    result["manifest_version"] = "pilot-1"
    return result


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def command_version(command: str) -> str | None:
    path = shutil.which(command)
    if not path:
        return None
    try:
        return subprocess.run([path, "--version"], capture_output=True, text=True, timeout=10).stdout.strip()[:500]
    except Exception as exc:  # pragma: no cover
        return f"version query failed: {exc}"


def main() -> None:
    for folder in [FIXTURES, MANIFESTS, LOGS, PILOT / "outputs" / "word", PILOT / "outputs" / "libreoffice", PILOT / "outputs" / "google_docs", PILOT / "source_semantics", PILOT / "destination_semantics", PILOT / "diffs", PILOT / "validators" / "pac", PILOT / "validators" / "acrobat", PILOT / "validators" / "verapdf", PILOT / "adjudication", PILOT / "screenshots", PILOT / "reports"]:
        folder.mkdir(parents=True, exist_ok=True)
    chart = FIXTURES / "F02_alt_text_chart.png"
    make_chart(chart)
    builders = [
        ("F01_HEADINGS", "A11yPreserve Heading Fixture", make_headings),
        ("F02_ALT_TEXT", "A11yPreserve Alternative Text Fixture", lambda: make_alt_text(chart)),
        ("F03_LISTS", "A11yPreserve List Fixture", make_lists),
        ("F04_TABLE", "A11yPreserve Table Fixture", make_table),
    ]
    hashes: dict[str, Any] = {}
    for fixture_id, title, builder in builders:
        document, features = builder()
        source = FIXTURES / f"{fixture_id}.docx"
        document.save(source)
        ground_truth = manifest(fixture_id, title, features)
        manifest_path = MANIFESTS / f"{fixture_id}.json"
        manifest_path.write_text(json.dumps(ground_truth, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
        hashes[fixture_id] = {"source_docx": sha256(source), "manifest": sha256(manifest_path)}
    (MANIFESTS / "hashes.json").write_text(json.dumps(hashes, indent=2) + "\n", encoding="utf-8")
    environment = {
        "recorded_utc": __import__("datetime").datetime.now(__import__("datetime").timezone.utc).isoformat(),
        "os": platform.platform(),
        "python": sys.version,
        "python_executable": sys.executable,
        "packages": {},
        "commands_on_path": {name: shutil.which(name) for name in ["soffice", "libreoffice", "verapdf", "pac", "pdftotext"]},
    }
    for package in ["docx", "pypdf", "PIL"]:
        try:
            module = __import__(package)
            environment["packages"][package] = getattr(module, "__version__", "installed")
        except Exception as exc:
            environment["packages"][package] = f"unavailable: {exc}"
    (LOGS / "environment_inventory.json").write_text(json.dumps(environment, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"fixtures": hashes, "environment": environment}, indent=2))


if __name__ == "__main__":
    main()
