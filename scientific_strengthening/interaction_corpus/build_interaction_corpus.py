"""Build six additive interaction fixtures with explicit source contracts.

The existing Core, Variant B, and integrated-study files are not touched.
Each document combines semantics that are independently covered elsewhere but
co-occurring here, preserving a deterministic source oracle.
"""

from __future__ import annotations

import hashlib
import importlib.util
import json
from pathlib import Path


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "scientific_strengthening" / "interaction_corpus"
SOURCES = OUT / "sources"

spec = importlib.util.spec_from_file_location("integrated_helpers", ROOT / "scripts" / "build_integrated_study.py")
helpers = importlib.util.module_from_spec(spec)
assert spec.loader is not None
spec.loader.exec_module(helpers)

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Inches  # noqa: E402


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def finish(document, title: str):
    document.core_properties.title = title
    document.core_properties.subject = "A11yPreserve interaction-corpus source contract"
    return document


def add_ordered_list(document, items):
    """Create an explicit decimal list (the helper's numId=1 is bullets)."""
    for text, level in items:
        paragraph = document.add_paragraph(style="List Number")
        paragraph.paragraph_format.left_indent = Inches(0.25 * level)
        run = paragraph.add_run(text)
        helpers.set_run_language(run, "en-US")
        ppr = paragraph._p.get_or_add_pPr()
        num_pr = ppr.find(helpers.q("w:numPr"))
        if num_pr is None:
            num_pr = helpers.OxmlElement("w:numPr")
            ppr.append(num_pr)
        ilvl = helpers.OxmlElement("w:ilvl")
        ilvl.set(helpers.q("w:val"), str(level))
        num_id = helpers.OxmlElement("w:numId")
        num_id.set(helpers.q("w:val"), "5")
        num_pr.append(ilvl)
        num_pr.append(num_id)


def ic01():
    title = "Interaction 01: Service navigation brief"
    d = finish(helpers.setup_document(title), title)
    helpers.add_title(d, title)
    helpers.add_body(d, "An authored brief that combines document language, navigable headings, and a named external link.")
    helpers.add_heading(d, "Purpose", 1)
    helpers.add_body(d, "This brief explains how readers can find a service request and where to ask for an accessible format.")
    helpers.add_heading(d, "Service steps", 1)
    helpers.add_heading(d, "Find the request", 2)
    helpers.add_body(d, "Start with the request title and read the short summary before opening the linked guidance.")
    helpers.add_heading(d, "Ask for help", 2)
    helpers.add_language_body(d, "The request form includes the Spanish phrase ", "información pública", " as a marked language shift.", "es-MX")
    helpers.add_link_paragraph(d, "The complete guidance is in the ", "accessibility standards overview", "https://www.w3.org/WAI/standards-guidelines/wcag/")
    return d, {
        "purpose": "heading hierarchy + hyperlink + inline language coexistence",
        "features": ["headings", "hyperlinks", "inline_language", "document_language"],
        "contracts": [
            {"feature": "headings", "expected": {"levels": [1, 1, 2, 2]}},
            {"feature": "hyperlinks", "expected": {"text": "accessibility standards overview", "target": "https://www.w3.org/WAI/standards-guidelines/wcag/"}},
            {"feature": "inline_language", "expected": {"text": "información pública", "language": "es-MX"}},
            {"feature": "document_language", "expected": {"language": "en-US"}},
        ],
    }


def ic02():
    title = "Interaction 02: Multilingual review checklist"
    d = finish(helpers.setup_document(title), title)
    helpers.add_title(d, title)
    helpers.add_heading(d, "Review goals", 1)
    helpers.add_body(d, "The checklist records a review sequence while preserving a short language shift inside explanatory prose.")
    helpers.add_language_body(d, "The reviewer may use the phrase ", "Buenos días", " when opening the session.", "es-MX")
    helpers.add_heading(d, "Checklist", 1)
    add_ordered_list(d, [
        ("Identify the audience.", 0),
        ("Check the document title.", 0),
        ("Check headings and reading order.", 1),
        ("Check the table header row.", 1),
        ("Record unresolved evidence.", 0),
    ])
    helpers.add_heading(d, "Decision", 1)
    helpers.add_body(d, "A reviewer records what the source contract establishes and does not infer equivalence from visual similarity alone.")
    return d, {
        "purpose": "nested ordered list + inline language + explanatory hierarchy",
        "features": ["headings", "nested_lists", "inline_language", "document_language"],
        "contracts": [
            {"feature": "headings", "expected": {"levels": [1, 1, 1]}},
            {"feature": "lists", "expected": {"depths": [0, 0, 1, 1, 0], "top_level_type": "ordered"}},
            {"feature": "inline_language", "expected": {"text": "Buenos días", "language": "es-MX"}},
            {"feature": "document_language", "expected": {"language": "en-US"}},
        ],
    }


def ic03():
    title = "Interaction 03: Evidence figure with caption and link"
    d = finish(helpers.setup_document(title), title)
    helpers.add_title(d, title)
    helpers.add_heading(d, "Evidence summary", 1)
    helpers.add_body(d, "The figure is part of the explanation, and the caption states the intended relationship between the visual and the surrounding text.")
    helpers.add_heading(d, "Visual finding", 2)
    helpers.add_image(d, "A rising line showing completed requests increasing across two reporting periods.", "Request trend")
    helpers.add_caption(d, "Figure 1. Completed requests increased across the reporting periods.")
    helpers.add_link_paragraph(d, "The interpretation follows the ", "public accessibility guidance", "https://www.w3.org/WAI/standards-guidelines/wcag/")
    helpers.add_heading(d, "Interpretation", 1)
    helpers.add_body(d, "The image alternative, caption, heading, and link are intentionally adjacent so that their relationships coexist in one document structure.")
    return d, {
        "purpose": "heading + informative image alternative + caption + hyperlink coexistence",
        "features": ["headings", "image_alt_text", "captions", "hyperlinks"],
        "contracts": [
            {"feature": "headings", "expected": {"levels": [1, 2, 1]}},
            {"feature": "image_alt_text", "expected": {"alt": "A rising line showing completed requests increasing across two reporting periods."}},
            {"feature": "captions", "expected": {"text": "Figure 1. Completed requests increased across the reporting periods."}},
            {"feature": "hyperlinks", "expected": {"text": "public accessibility guidance", "target": "https://www.w3.org/WAI/standards-guidelines/wcag/"}},
        ],
    }


def ic04():
    title = "Interaction 04: Service measures table"
    d = finish(helpers.setup_document(title), title)
    helpers.add_title(d, title)
    helpers.add_heading(d, "Measures and actions", 1)
    helpers.add_body(d, "The table reports measures, while the list below records the actions that explain how the measures are used.")
    helpers.add_table(d, ["Measure", "Current", "Next step"], [["Requests", "148", "Review"], ["Wait days", "5", "Maintain"], ["Follow-ups", "12", "Explain"]], "Table 1. Service measures and next steps.")
    helpers.add_heading(d, "Action list", 2)
    helpers.add_list(d, [("Review the request count.", 0), ("Explain the next step.", 1), ("Record the result.", 0)])
    helpers.add_heading(d, "Reading note", 1)
    helpers.add_body(d, "The header row, list nesting, and heading hierarchy are separate source properties that must coexist after conversion.")
    return d, {
        "purpose": "table headers + list nesting + heading hierarchy",
        "features": ["headings", "table_headers", "lists"],
        "contracts": [
            {"feature": "headings", "expected": {"levels": [1, 2, 1]}},
            {"feature": "table", "expected": {"headers": ["Measure", "Current", "Next step"], "rows": 4}},
            {"feature": "lists", "expected": {"depths": [0, 1, 0]}},
        ],
    }


def ic05():
    title = "Interaction 05: Linked source note"
    d = finish(helpers.setup_document(title), title)
    helpers.add_title(d, title)
    helpers.add_heading(d, "Source note", 1)
    helpers.add_body(d, "This short note identifies the evidence used for a public-information example[[FN]].")
    helpers.add_link_paragraph(d, "The related standard is the ", "accessibility standards overview", "https://www.w3.org/WAI/standards-guidelines/wcag/")
    helpers.add_heading(d, "Use", 2)
    helpers.add_body(d, "Readers can navigate from the heading to the linked source and the footnote reference without relying on placement alone.")
    raw = OUT / "_IC05_footnote_raw.docx"
    final = SOURCES / "IC05_LINKED_SOURCE_NOTE.docx"
    d.save(raw)
    helpers.insert_note(str(raw), str(final), "footnote", "[[FN]]", "This note identifies the source used for the public-information example.")
    raw.unlink(missing_ok=True)
    return Document(str(final)), {
        "purpose": "footnote reference + hyperlink + heading coexistence",
        "features": ["headings", "footnotes", "hyperlinks"],
        "contracts": [
            {"feature": "headings", "expected": {"levels": [1, 2]}},
            {"feature": "footnotes", "expected": {"text": "This note identifies the source used for the public-information example."}},
            {"feature": "hyperlinks", "expected": {"text": "accessibility standards overview", "target": "https://www.w3.org/WAI/standards-guidelines/wcag/"}},
        ],
    }


def ic06():
    title = "Interaction 06: Accessible service mini-report"
    d = finish(helpers.setup_document(title), title)
    helpers.add_title(d, title)
    helpers.add_body(d, "A multi-section mini-report combining narrative, hierarchy, lists, language, a table, a figure, a caption, and a link.")
    helpers.add_heading(d, "Overview", 1)
    helpers.add_body(d, "The report describes a service update and records the structures that carry its meaning across sections.")
    helpers.add_heading(d, "Changes", 1)
    helpers.add_heading(d, "What improved", 2)
    helpers.add_list(d, [("Shorter intake instructions.", 0), ("Clearer contact route.", 0), ("Published format request process.", 1)])
    helpers.add_language_body(d, "The example request is labeled ", "información pública", " in Spanish.", "es-MX")
    helpers.add_heading(d, "Measures", 1)
    helpers.add_table(d, ["Measure", "Before", "After"], [["Intake steps", "6", "4"], ["Wait days", "8", "5"]], "Table 1. Changes in service measures.")
    helpers.add_heading(d, "Evidence", 1)
    helpers.add_image(d, "A simple process diagram showing intake, review, and response stages.", "Service process")
    helpers.add_caption(d, "Figure 1. The service process has intake, review, and response stages.")
    helpers.add_link_paragraph(d, "The publication checklist follows the ", "Web Content Accessibility Guidelines", "https://www.w3.org/WAI/standards-guidelines/wcag/")
    return d, {
        "purpose": "multi-page-style mini-report with cross-feature coexistence",
        "features": ["headings", "lists", "table_headers", "image_alt_text", "captions", "inline_language", "hyperlinks"],
        "contracts": [
            {"feature": "headings", "expected": {"levels": [1, 1, 2, 1, 1]}},
            {"feature": "lists", "expected": {"depths": [0, 0, 1]}},
            {"feature": "table", "expected": {"headers": ["Measure", "Before", "After"], "rows": 3}},
            {"feature": "image_alt_text", "expected": {"alt": "A simple process diagram showing intake, review, and response stages."}},
            {"feature": "captions", "expected": {"text": "Figure 1. The service process has intake, review, and response stages."}},
            {"feature": "inline_language", "expected": {"text": "información pública", "language": "es-MX"}},
            {"feature": "hyperlinks", "expected": {"text": "Web Content Accessibility Guidelines", "target": "https://www.w3.org/WAI/standards-guidelines/wcag/"}},
        ],
    }


def main() -> None:
    SOURCES.mkdir(parents=True, exist_ok=True)
    builders = [ic01, ic02, ic03, ic04, ic05, ic06]
    records = []
    for index, builder in enumerate(builders, start=1):
        document, metadata = builder()
        document_id = f"IC{index:02d}"
        path = next((p for p in SOURCES.glob(f"{document_id}_*.docx")), None)
        if path is None:
            path = SOURCES / f"{document_id}_{metadata['purpose'].split()[0].upper()}.docx"
        document.save(path)
        records.append({
            "document_id": document_id,
            "source_docx": str(path.relative_to(ROOT)).replace("\\", "/"),
            "source_sha256": sha256(path),
            **metadata,
        })
    manifest = {"corpus": "interaction_corpus", "design": "controlled feature coexistence; not a random real-world sample", "documents": records}
    (OUT / "INTERACTION_CONTRACTS.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps({"documents": len(records), "contracts": sum(len(r["contracts"]) for r in records)}, indent=2))


if __name__ == "__main__":
    main()
